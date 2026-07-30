import io
import json
import hashlib
import threading
from datetime import datetime
from flask import Blueprint, request, jsonify, g, send_file, current_app
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from ..extensions import db, socketio
from ..models.user import UserRole
from ..models.student import Student
from ..models.placement import JobRole, JobRoleSkill, ResumeScan, ScanStatus, CompanyCriteria, LearningAssignment
from ..utils.auth_middleware import require_auth, require_hod, require_faculty_or_hod
from ..utils.cloudinary_helper import upload_file
from ..utils.groq_helper import analyze_resume_with_groq

placement_bp = Blueprint("placement", __name__)

DEFAULT_ROLES = [
    {
        "name": "Full-Stack Developer",
        "description": "Design and build responsive web applications across front-end and back-end architectures.",
        "skills": [
            {"skill_name": "React", "weight": 2, "is_core": True},
            {"skill_name": "Node.js", "weight": 2, "is_core": True},
            {"skill_name": "Python", "weight": 1, "is_core": True},
            {"skill_name": "SQL", "weight": 2, "is_core": True},
            {"skill_name": "REST API", "weight": 2, "is_core": True},
            {"skill_name": "Git", "weight": 1, "is_core": True},
            {"skill_name": "Docker", "weight": 1, "is_core": False},
            {"skill_name": "System Design", "weight": 1, "is_core": False},
            {"skill_name": "TypeScript", "weight": 1, "is_core": False},
            {"skill_name": "HTML/CSS", "weight": 1, "is_core": True},
        ]
    },
    {
        "name": "Data Scientist / AI Engineer",
        "description": "Extract insights from complex data models and deploy artificial intelligence/machine learning solutions.",
        "skills": [
            {"skill_name": "Python", "weight": 2, "is_core": True},
            {"skill_name": "Machine Learning", "weight": 2, "is_core": True},
            {"skill_name": "PyTorch / TensorFlow", "weight": 1, "is_core": False},
            {"skill_name": "SQL", "weight": 2, "is_core": True},
            {"skill_name": "Pandas / NumPy", "weight": 2, "is_core": True},
            {"skill_name": "Data Visualization", "weight": 1, "is_core": True},
            {"skill_name": "Deep Learning", "weight": 1, "is_core": False},
            {"skill_name": "Statistics", "weight": 1, "is_core": True},
            {"skill_name": "Scikit-Learn", "weight": 1, "is_core": True},
            {"skill_name": "Git", "weight": 1, "is_core": True},
        ]
    },
    {
        "name": "Cloud & DevOps Engineer",
        "description": "Automate deployment pipelines and maintain robust cloud infrastructure & microservices.",
        "skills": [
            {"skill_name": "Linux / Bash", "weight": 2, "is_core": True},
            {"skill_name": "Docker", "weight": 2, "is_core": True},
            {"skill_name": "Kubernetes", "weight": 1, "is_core": False},
            {"skill_name": "AWS / Cloud", "weight": 2, "is_core": True},
            {"skill_name": "CI/CD Pipelines", "weight": 2, "is_core": True},
            {"skill_name": "Git", "weight": 1, "is_core": True},
            {"skill_name": "Networking & Security", "weight": 1, "is_core": True},
            {"skill_name": "Python / Shell Scripting", "weight": 1, "is_core": True},
            {"skill_name": "Terraform", "weight": 1, "is_core": False},
        ]
    }
]


def _seed_roles_if_needed():
    """Ensure default JobRole entries exist."""
    try:
        if JobRole.query.count() == 0:
            for r_data in DEFAULT_ROLES:
                role = JobRole(name=r_data["name"], description=r_data["description"])
                db.session.add(role)
                db.session.flush()
                for s_data in r_data["skills"]:
                    skill = JobRoleSkill(
                        role_id=role.id,
                        skill_name=s_data["skill_name"],
                        weight=s_data["weight"],
                        is_core=s_data["is_core"]
                    )
                    db.session.add(skill)
            db.session.commit()
    except Exception as e:
        db.session.rollback()
        current_app.logger.warning(f"Error seeding job roles: {e}")


def _extract_pdf_text(file_bytes: bytes) -> str:
    """Extract plain text from PDF in memory using pdfplumber."""
    extracted_text = ""
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    extracted_text += t + "\n"
    except Exception as e:
        current_app.logger.error(f"Error extracting PDF text: {e}")
    return extracted_text.strip()


def _process_resume_in_background(app, scan_id, user_id, resume_text, role_name, required_skills):
    """Background worker function executing Groq AI evaluation."""
    with app.app_context():
        scan = ResumeScan.query.get(scan_id)
        if not scan:
            return

        try:
            scan.status = ScanStatus.PROCESSING
            db.session.commit()

            ai_res = analyze_resume_with_groq(resume_text, role_name, required_skills)

            scan.match_score = ai_res["match_score"]
            scan.matched_skills_json = json.dumps(ai_res["matched_skills"])
            scan.missing_skills_json = json.dumps(ai_res["missing_skills"])
            scan.action_plan_json = json.dumps(ai_res["action_plan"])
            scan.ai_raw_response = ai_res.get("raw_response", "")
            scan.status = ScanStatus.COMPLETE
            scan.completed_at = datetime.utcnow()
            db.session.commit()

            # Broadcast realtime socket event
            try:
                socketio.emit("scan_complete", {
                    "scan_id": scan.id,
                    "match_score": scan.match_score,
                    "status": "COMPLETE"
                }, room=f"user_{user_id}")
            except Exception as se:
                current_app.logger.warning(f"Socket emit notice warning: {se}")

        except Exception as e:
            current_app.logger.error(f"Resume scan id {scan_id} background evaluation error: {e}")
            db.session.rollback()
            scan.status = ScanStatus.FAILED
            scan.error_message = str(e)
            db.session.commit()

            try:
                socketio.emit("scan_failed", {
                    "scan_id": scan.id,
                    "error": str(e),
                    "status": "FAILED"
                }, room=f"user_{user_id}")
            except Exception:
                pass


# ─── Routes ─────────────────────────────────────────────────────────────

@placement_bp.route("/roles", methods=["GET"])
@require_auth
def get_job_roles():
    _seed_roles_if_needed()
    roles = JobRole.query.filter_by(is_active=True).all()
    return jsonify({"roles": [r.to_dict() for r in roles]})


@placement_bp.route("/scan", methods=["POST"])
@require_auth
def submit_resume_scan():
    user = g.current_user
    if user.role != UserRole.STUDENT or not user.student_profile:
        return jsonify({"error": "Only student profiles can evaluate resume placement readiness."}), 403

    student = user.student_profile
    role_id = request.form.get("role_id")
    if not role_id:
        return jsonify({"error": "Target job role ID is required."}), 400

    role = JobRole.query.get(role_id)
    if not role:
        return jsonify({"error": "Invalid job role ID."}), 404

    if "resume" not in request.files:
        return jsonify({"error": "Resume PDF file is required."}), 400

    file = request.files["resume"]
    if not file.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Only PDF resume documents are supported."}), 400

    file_bytes = file.read()
    if len(file_bytes) > 10 * 1024 * 1024:
        return jsonify({"error": "File size exceeds 10MB limit."}), 400

    file.seek(0)
    # Upload to Cloudinary
    c_res = upload_file(file, folder="placement_resumes")

    resume_text = _extract_pdf_text(file_bytes)
    if not resume_text or len(resume_text) < 50:
        return jsonify({"error": "Unable to extract readable text from PDF. Please upload a non-scanned text PDF."}), 400

    text_hash = hashlib.sha256(resume_text.encode("utf-8")).hexdigest()

    # Create ResumeScan record in PENDING state
    scan = ResumeScan(
        student_id=student.id,
        role_id=role.id,
        resume_url=c_res["url"],
        resume_public_id=c_res.get("public_id"),
        resume_text_hash=text_hash,
        status=ScanStatus.PENDING
    )
    db.session.add(scan)
    db.session.commit()

    # Prepare required skill names
    skills_list = [s.skill_name for s in role.required_skills]

    # Dispatch to background thread
    app = current_app._get_current_object()
    thread = threading.Thread(
        target=_process_resume_in_background,
        args=(app, scan.id, user.id, resume_text, role.name, skills_list)
    )
    thread.daemon = True
    thread.start()

    return jsonify({
        "message": "Resume uploaded successfully. Analysis in progress.",
        "scan_id": scan.id,
        "status": "PENDING"
    }), 202


@placement_bp.route("/scan/<int:scan_id>", methods=["GET"])
@require_auth
def get_scan_detail(scan_id):
    user = g.current_user
    scan = ResumeScan.query.get_or_404(scan_id)

    # Permission check: Student can view own scan, Faculty/HOD can view
    if user.role == UserRole.STUDENT and user.student_profile:
        if scan.student_id != user.student_profile.id:
            return jsonify({"error": "Access denied"}), 403

    return jsonify({"scan": scan.to_dict()})


@placement_bp.route("/scan/latest", methods=["GET"])
@require_auth
def get_latest_student_scan():
    user = g.current_user
    if user.role != UserRole.STUDENT or not user.student_profile:
        return jsonify({"error": "Only student profiles have scan records."}), 403

    role_id = request.args.get("role_id")
    query = ResumeScan.query.filter_by(student_id=user.student_profile.id)
    if role_id:
        query = query.filter_by(role_id=role_id)

    scan = query.order_by(ResumeScan.id.desc()).first()
    return jsonify({"scan": scan.to_dict() if scan else None})


@placement_bp.route("/scans", methods=["GET"])
@require_auth
def get_student_scans():
    user = g.current_user
    if user.role != UserRole.STUDENT or not user.student_profile:
        return jsonify({"scans": []})

    scans = ResumeScan.query.filter_by(student_id=user.student_profile.id).order_by(ResumeScan.id.desc()).all()
    return jsonify({"scans": [s.to_dict() for s in scans]})


@placement_bp.route("/faculty/summary", methods=["GET"])
@require_faculty_or_hod
def get_faculty_class_summary():
    user = g.current_user
    class_name = request.args.get("class_name")

    if user.role == UserRole.FACULTY:
        fp = user.faculty_profile
        if not fp or not fp.class_incharge_of:
            return jsonify({"error": "Only Class Incharges and HOD can view placement analytics."}), 403
        # Class Incharge can only view their own assigned class
        class_name = fp.class_incharge_of

    if not class_name:
        class_name = "UG_3B"

    # Query all students in target class
    students = Student.query.filter_by(class_name=class_name).all()
    student_ids = [s.id for s in students]

    if not student_ids:
        return jsonify({
            "class_name": class_name,
            "total_students": 0,
            "job_ready_count": 0,
            "moderate_count": 0,
            "needs_support_count": 0,
            "top_skill_gaps": [],
            "students": []
        })

    # Fetch latest scan for each student
    latest_scans = []
    skill_gap_counts = {}

    for s_id in student_ids:
        scan = ResumeScan.query.filter_by(student_id=s_id, status=ScanStatus.COMPLETE).order_by(ResumeScan.id.desc()).first()
        if scan:
            latest_scans.append(scan)
            for skill in scan.get_missing_skills():
                skill_gap_counts[skill] = skill_gap_counts.get(skill, 0) + 1

    job_ready = sum(1 for s in latest_scans if s.match_score and s.match_score >= 70)
    moderate = sum(1 for s in latest_scans if s.match_score and 40 <= s.match_score < 70)
    needs_support = sum(1 for s in latest_scans if s.match_score and s.match_score < 40)

    # Sort top missing skill gaps
    sorted_gaps = [{"skill_name": k, "count": v} for k, v in sorted(skill_gap_counts.items(), key=lambda x: x[1], reverse=True)]

    # Serialize student records with scan status
    student_records = []
    for s in students:
        s_scan = next((sc for sc in latest_scans if sc.student_id == s.id), None)
        student_records.append({
            "student_id": s.id,
            "name": s.user.name,
            "roll_number": s.roll_number,
            "profile_picture": s.user.profile_picture,
            "has_scan": s_scan is not None,
            "latest_score": s_scan.match_score if s_scan else None,
            "target_role": s_scan.role.name if s_scan and s_scan.role else None,
            "missing_skills": s_scan.get_missing_skills() if s_scan else []
        })

    return jsonify({
        "class_name": class_name,
        "total_students": len(students),
        "evaluated_students": len(latest_scans),
        "job_ready_count": job_ready,
        "moderate_count": moderate,
        "needs_support_count": needs_support,
        "top_skill_gaps": sorted_gaps,
        "students": student_records
    })


@placement_bp.route("/faculty/assign", methods=["POST"])
@require_faculty_or_hod
def assign_learning_resource():
    user = g.current_user
    data = request.json or {}

    skill_name = data.get("skill_name")
    resource_title = data.get("resource_title")

    if not skill_name or not resource_title:
        return jsonify({"error": "Skill name and resource title are required."}), 400

    assignment = LearningAssignment(
        assigned_by=user.id,
        class_name=data.get("class_name"),
        student_id=data.get("student_id"),
        skill_name=skill_name,
        resource_title=resource_title,
        resource_url=data.get("resource_url", ""),
        resource_type=data.get("resource_type", "course")
    )

    db.session.add(assignment)
    db.session.commit()

    return jsonify({"message": "Learning resource assigned successfully.", "assignment": assignment.to_dict()})


@placement_bp.route("/faculty/assignments", methods=["GET"])
@require_auth
def get_learning_assignments():
    user = g.current_user
    class_name = request.args.get("class_name")

    query = LearningAssignment.query
    if user.role == UserRole.STUDENT and user.student_profile:
        query = query.filter(
            (LearningAssignment.student_id == user.student_profile.id) |
            (LearningAssignment.class_name == user.student_profile.class_name)
        )
    elif class_name:
        query = query.filter_by(class_name=class_name)

    assignments = query.order_by(LearningAssignment.id.desc()).all()
    return jsonify({"assignments": [a.to_dict() for a in assignments]})


@placement_bp.route("/hod/matrix", methods=["GET"])
@require_hod
def get_hod_placement_matrix():
    classes_list = ["UG_1A", "UG_1B", "UG_2A", "UG_2B", "UG_3A", "UG_3B", "PG_1A", "PG_2A"]
    matrix = []

    for c in classes_list:
        students = Student.query.filter_by(class_name=c).all()
        s_ids = [s.id for s in students]

        if not s_ids:
            matrix.append({
                "class_name": c,
                "total_students": 0,
                "evaluated_count": 0,
                "job_ready_count": 0,
                "avg_match_score": 0
            })
            continue

        scans = ResumeScan.query.filter(
            ResumeScan.student_id.in_(s_ids),
            ResumeScan.status == ScanStatus.COMPLETE
        ).all()

        latest_per_student = {}
        for sc in scans:
            if sc.student_id not in latest_per_student or sc.id > latest_per_student[sc.student_id].id:
                latest_per_student[sc.student_id] = sc

        scores = [sc.match_score for sc in latest_per_student.values() if sc.match_score is not None]
        job_ready = sum(1 for sc in latest_per_student.values() if sc.match_score and sc.match_score >= 70)
        avg_score = round(sum(scores) / len(scores), 1) if scores else 0

        matrix.append({
            "class_name": c,
            "total_students": len(students),
            "evaluated_count": len(latest_per_student),
            "job_ready_count": job_ready,
            "avg_match_score": avg_score
        })

    return jsonify({"matrix": matrix})


@placement_bp.route("/hod/company-criteria", methods=["GET", "POST"])
@require_hod
def handle_company_criteria():
    user = g.current_user
    if request.method == "POST":
        data = request.json or {}
        company_name = data.get("company_name")
        role_id = data.get("role_id")

        if not company_name or not role_id:
            return jsonify({"error": "Company name and role ID are required."}), 400

        criteria = CompanyCriteria(
            company_name=company_name,
            role_id=role_id,
            min_match_score=int(data.get("min_match_score", 70)),
            min_cgpa=str(data.get("min_cgpa", "6.5")),
            required_skills_json=json.dumps(data.get("required_skills", [])),
            created_by=user.id
        )
        db.session.add(criteria)
        db.session.commit()
        return jsonify({"message": "Company criteria created.", "criteria": criteria.to_dict()})

    criterias = CompanyCriteria.query.order_by(CompanyCriteria.id.desc()).all()
    return jsonify({"criterias": [c.to_dict() for c in criterias]})


@placement_bp.route("/hod/company-criteria/<int:criteria_id>/eligible-students", methods=["GET"])
@require_hod
def get_eligible_students(criteria_id):
    criteria = CompanyCriteria.query.get_or_404(criteria_id)
    students = Student.query.all()
    eligible = []

    for s in students:
        scan = ResumeScan.query.filter_by(
            student_id=s.id,
            role_id=criteria.role_id,
            status=ScanStatus.COMPLETE
        ).order_by(ResumeScan.id.desc()).first()

        if scan and scan.match_score and scan.match_score >= criteria.min_match_score:
            eligible.append({
                "student_id": s.id,
                "name": s.user.name,
                "email": s.user.email,
                "roll_number": s.roll_number,
                "class_name": s.class_name,
                "profile_picture": s.user.profile_picture,
                "match_score": scan.match_score,
                "matched_skills": scan.get_matched_skills(),
                "resume_url": scan.resume_url
            })

    return jsonify({
        "criteria": criteria.to_dict(),
        "eligible_students_count": len(eligible),
        "students": eligible
    })


@placement_bp.route("/hod/export", methods=["GET"])
@require_hod
def export_placement_report():
    wb = openpyxl.Workbook()

    # Sheet 1: Matrix Summary
    ws_matrix = wb.active
    ws_matrix.title = "Placement Matrix"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="F05A2B", end_color="F05A2B", fill_type="solid")

    headers_matrix = ["Class Name", "Total Students", "Evaluated Students", "Job-Ready (≥70%)", "Avg Match Score"]
    ws_matrix.append(headers_matrix)

    for col_num in range(1, len(headers_matrix) + 1):
        cell = ws_matrix.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    classes_list = ["UG_1A", "UG_1B", "UG_2A", "UG_2B", "UG_3A", "UG_3B", "PG_1A", "PG_2A"]
    for c in classes_list:
        students = Student.query.filter_by(class_name=c).all()
        s_ids = [s.id for s in students]
        if not s_ids:
            ws_matrix.append([c.replace("_", " "), 0, 0, 0, "0%"])
            continue

        scans = ResumeScan.query.filter(
            ResumeScan.student_id.in_(s_ids),
            ResumeScan.status == ScanStatus.COMPLETE
        ).all()

        latest = {}
        for sc in scans:
            if sc.student_id not in latest or sc.id > latest[sc.student_id].id:
                latest[sc.student_id] = sc

        scores = [sc.match_score for sc in latest.values() if sc.match_score is not None]
        job_ready = sum(1 for sc in latest.values() if sc.match_score and sc.match_score >= 70)
        avg_score = round(sum(scores) / len(scores), 1) if scores else 0

        ws_matrix.append([c.replace("_", " "), len(students), len(latest), job_ready, f"{avg_score}%"])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"Hive_Placement_Readiness_Report_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
